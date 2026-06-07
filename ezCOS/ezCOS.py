# -*- coding: utf-8 -*-
#
# ezCOS.py -
#
# Authors: 
# Géry Casiez https://gery.casiez.net
#
# 2026
# 
# BSD License https://opensource.org/licenses/BSD-3-Clause
#
# Redistribution and use in source and binary forms, with or without
# modification, are permitted provided that the following conditions are met:
#
#  1. Redistributions of source code must retain the above copyright notice, this list of conditions
# and the following disclaimer.
#
# 2. Redistributions in binary form must reproduce the above copyright notice, this list of conditions
# and the following disclaimer in the documentation and/or other materials provided with the distribution.
# 
# 3. Neither the name of the copyright holder nor the names of its contributors may be used to endorse or
# promote products derived from this software without specific prior written permission.

# THIS SOFTWARE IS PROVIDED BY THE COPYRIGHT HOLDERS AND CONTRIBUTORS "AS IS" AND ANY EXPRESS OR IMPLIED WARRANTIES, 
# INCLUDING, BUT NOT LIMITED TO, THE IMPLIED WARRANTIES OF MERCHANTABILITY AND FITNESS FOR A PARTICULAR PURPOSE ARE
# DISCLAIMED. IN NO EVENT SHALL THE COPYRIGHT HOLDER OR CONTRIBUTORS BE LIABLE FOR ANY DIRECT, INDIRECT, INCIDENTAL,
# SPECIAL, EXEMPLARY, OR CONSEQUENTIAL DAMAGES (INCLUDING, BUT NOT LIMITED TO, PROCUREMENT OF SUBSTITUTE GOODS OR SERVICES;
# LOSS OF USE, DATA, OR PROFITS; OR BUSINESS INTERRUPTION) HOWEVER CAUSED AND ON ANY THEORY OF LIABILITY, WHETHER IN 
# CONTRACT, STRICT LIABILITY, OR TORT (INCLUDING NEGLIGENCE OR OTHERWISE) ARISING IN ANY WAY OUT OF THE USE OF THIS
# SOFTWARE, EVEN IF ADVISED OF THE POSSIBILITY OF SUCH DAMAGE.
#

import argparse
import configparser
from datetime import datetime
import os
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import quote_sheetname
from odysseeapi.OdysseeCOS import OdysseeCOS
import zipfile
import pandas as pd
import sys
from tqdm import tqdm
from unidecode import unidecode
import tqdm

from .utils.tools import getColumnInfo, populateSheet, computeAge, loadJson, saveJson, createDir, getTextFromPDF, extractInfo, formatPrenom, formatLieu, formattxt, cleantext, thunderbird, trAvis, extractInfoRapport


def populateXLSX(template_dir, template_file, odyssee, candidates_folder):
    wb = load_workbook(filename=f"{template_dir}/{template_file}", data_only=False)
    ws = wb["MembresComite"]

    # Membres du comité de sélection
    membres = odyssee.get_committee_members()
    etablissements = odyssee.get_institutions()

    mapping = {
        "idKeycloak": {"column": "ID"},
        "nomUsage": {"column": "Nom"},
        "prenom": {"column": "Prénom", "transformation": formatPrenom},
        "civilite": {"column": "H/F", "transformation": lambda x: "H" if x == 1 else "F" },
        "rang": {"column": "Corps"},
        "etablissement": {"column": "Etablissement", "transformation": lambda x: next((formatLieu(e['libelle']) for e in etablissements if e['id'] == x), "") },
    }
    populateSheet(ws, membres, mapping)

    # Candidats
    candidats = odyssee.get_candidates()
    candidatsDetailed = odyssee.get_candidates_with_details()

    candidates_infos_fromPDF = getCandidatesInfos(candidates_folder)

    # Ajout de l'ID dans les détails des candidats pour faciliter le mapping
    for c in candidatsDetailed:
        # Ajout de l'ID Keycloak dans les détails des candidats pour faciliter le mapping
        idKeycloak = next((cand['candidat']['idKeycloakActuel'] for cand in candidats if cand['candidat']['nomUsage'] == c['informationsGeneralesCandidatureDto']['nomUsage'] and cand['candidat']['prenom'] == c['informationsGeneralesCandidatureDto']['prenom']), None)
        c['informationsGeneralesCandidatureDto']['idKeycloak'] = idKeycloak

        # Lieu de soutenance
        lieu = c['candidatureThese']['lieuSoutenanceEtablissement']
        if lieu is None:
            lieu = c['candidatureThese']['lieuSoutenanceAutre']
        else:
            lieu = next((formatLieu(e['libelle']) for e in etablissements if e['id'] == lieu), "")
        c['candidatureThese']['lieuSoutenance'] = lieu

        # Direction de thèse
        c['candidatureThese']['directeurThese'] = f"{formatPrenom(c['candidatureThese']['directeur']['prenom'])} {c['candidatureThese']['directeur']['nom'].upper()}"

        # Jury de thèse
        membresJury = []
        if c['candidatureJuryThese']['president'] is not None:
            membresJury.append(f"{formatPrenom(c['candidatureJuryThese']['president']['prenom'])} {c['candidatureJuryThese']['president']['nom'].upper()} (Président)")
        for membre in c['candidatureJuryThese']['membres']:
            membresJury.append(f"{formatPrenom(membre['prenom'])} {membre['nom'].upper()}")
        c['candidatureThese']['juryThese'] = ", ".join(membresJury)

        # Mots clés recherche
        keywords = odyssee.get_keywords()
        motsCles = []
        for mc in c['candidatureBlocMotCle']['motsCles']:
            if mc['motCleLibre'] is not None:
                motsCles.append(mc['motCleLibre'])
            else:
                mcle = next((k['libelle'] for k in keywords if k['id'] == mc['motCle']), "")
                motsCles.append(mcle.split("- ")[1])
        c['candidatureBlocMotCle']['motsCles'] = "; ".join(motsCles)

        # Lieu d'exercice du candidat
        idlieu = c['candidatureLieuRecherche']['etablissementLieuRecherche']
        c['candidatureLieuRecherche']['etablissementLieuRecherche'] = next((formatLieu(e['libelle']) for e in etablissements if e['id'] == idlieu), "")

        # Ajout des infos extraites des PDF
        candidate_infos = next((ci for ci in candidates_infos_fromPDF if ci['id'] == c['informationsGeneralesCandidatureDto']['idKeycloak']), None)
        if candidate_infos:
            c['informationsGeneralesCandidatureDto']['nom'] = candidate_infos.get('Nom', None)
            c['informationsGeneralesCandidatureDto']['email'] = candidate_infos.get('Email', None)
            c['informationsGeneralesCandidatureDto']['dateNaissance'] = candidate_infos.get('Date de naissance', None)
            c['informationsGeneralesCandidatureDto']['lieuNaissance'] = candidate_infos.get('Lieu de naissance', None)
            c['informationsGeneralesCandidatureDto']['paysNationalite'] = candidate_infos.get('Pays de nationalité', None)
            c['informationsGeneralesCandidatureDto']['telephone'] = candidate_infos.get('Téléphone', None)
            c['informationsGeneralesCandidatureDto']['adressePostale'] = candidate_infos.get('Adresse postale', None)
            c['informationsGeneralesCandidatureDto']['complementAdresse'] = candidate_infos.get('Adresse comp.', None)
            c['informationsGeneralesCandidatureDto']['codePostal'] = candidate_infos.get('Code postal', None)
            c['informationsGeneralesCandidatureDto']['ville'] = candidate_infos.get('Ville', None)
            c['informationsGeneralesCandidatureDto']['pays'] = candidate_infos.get('Pays', None)

        # Data validation Interne / Externe
        columnsInfos = getColumnInfo(ws)
        dvInterneExterne = DataValidation(type="list", formula1='"Interne,Externe"', allow_blank=True)
        ws.add_data_validation(dvInterneExterne)
        for row_index in range(2, ws.max_row + 1):
            index_interne_externe = next((col['index'] for col in columnsInfos if col['column'] == "Interne / Externe"), None)
            dvInterneExterne.add(ws.cell(row=row_index, column=index_interne_externe + 1))


    ws = wb["Candidats"]
    mapping = {
        "informationsGeneralesCandidatureDto:idKeycloak": {"column": "ID"},
        "informationsGeneralesCandidatureDto:nomUsage": {"column": "Nom d'usage", "transformation": lambda x: x.upper() if x else None},
        "informationsGeneralesCandidatureDto:prenom": {"column": "Prénom", "transformation": formatPrenom},
        "informationsGeneralesCandidatureDto:civilite": {"column": "Civilité", "transformation": lambda x: "M." if x == 1 else "Mme" },
        "candidatureThese:titre": {"column": "Titre thèse"},
        "candidatureThese:dateObtention": {"column": "Date soutenance", "transformation": lambda x: datetime.strptime(x, "%Y-%m-%d").strftime("%d/%m/%Y") if x else None},
        "candidatureThese:lieuSoutenance": {"column": "Lieu soutenance"},
        "candidatureThese:directeurThese": {"column": "Directeur Thèse"},
        "candidatureThese:juryThese": {"column": "Jury"},
        "candidatureBlocMotCle:description": {"column": "Description recherche"},
        "candidatureBlocMotCle:motsCles": {"column": "Mots clés"},
        "candidatureLieuRecherche:etablissementLieuRecherche": {"column": "Lieu d'exercice"},
        "informationsGeneralesCandidatureDto:nom": {"column": "Nom"},
        "informationsGeneralesCandidatureDto:email": {"column": "Email"},
        "informationsGeneralesCandidatureDto:dateNaissance": {"column": "Date de naissance"},
        "informationsGeneralesCandidatureDto:lieuNaissance": {"column": "Lieu de naissance"},
        "informationsGeneralesCandidatureDto:paysNationalite": {"column": "Pays de nationalité"},
        "informationsGeneralesCandidatureDto:telephone": {"column": "Téléphone"},
        "informationsGeneralesCandidatureDto:adressePostale": {"column": "Adresse postale"},
        "informationsGeneralesCandidatureDto:complementAdresse": {"column": "Adresse comp."},
        "informationsGeneralesCandidatureDto:codePostal": {"column": "Code postal"},
        "informationsGeneralesCandidatureDto:ville": {"column": "Ville"},
        "informationsGeneralesCandidatureDto:pays": {"column": "Pays"}
    }
    populateSheet(ws, candidatsDetailed, mapping)

    # Validation des données
    columnsInfos = getColumnInfo(ws)
    dvAvis = DataValidation(type="list", formula1='"Très favorable,Favorable,Réservé,Défavorable"', allow_blank=True)
    ws.add_data_validation(dvAvis)

    dvRapporteur = DataValidation(type="list", formula1=f"{quote_sheetname('MembresComite')}!$K$2:$K${len(membres)+1}", allow_blank=True)
    ws.add_data_validation(dvRapporteur)

    for row_index in range(2, ws.max_row + 1):
        index_avis_rapp1 = next((col['index'] for col in columnsInfos if col['column'] == "AvisRapp1"), None)
        dvAvis.add(ws.cell(row=row_index, column=index_avis_rapp1 + 1))
        index_avis_rapp2 = next((col['index'] for col in columnsInfos if col['column'] == "AvisRapp2"), None)
        dvAvis.add(ws.cell(row=row_index, column=index_avis_rapp2 + 1))
        
        index_rapporteur1 = next((col['index'] for col in columnsInfos if col['column'] == "Rapporteur1"), None)
        dvRapporteur.add(ws.cell(row=row_index, column=index_rapporteur1 + 1))
        index_rapporteur2 = next((col['index'] for col in columnsInfos if col['column'] == "Rapporteur2"), None)
        dvRapporteur.add(ws.cell(row=row_index, column=index_rapporteur2 + 1))

    wb.save(f"comitePopulated.xlsx")

def unzip_applications(directory):
    zip_files = [f for f in os.listdir(directory) if f.endswith('.zip')]

    for zip_file in zip_files:
        path_to_zip_file = os.path.join(directory, zip_file)
        directory_to_extract_to = os.path.join(directory, os.path.splitext(zip_file)[0])
    
        with zipfile.ZipFile(path_to_zip_file, 'r') as zip_ref:
            zip_ref.extractall(directory_to_extract_to)

def extractInfos(infos, txt):
    res = {}
    for info in infos:
        i = extractInfo(info['odyssee'], txt, info.get('stopWord', None))
        if i and 'operation' in info:
            i = info['operation'](i)
        res[info['target']] = i
    return res

def getPdfFiles(folder):
    pdf_files = []
    for root, _, files in os.walk(folder):
        for file in files:
            if file.startswith("Edition_Dossier") and file.endswith(".pdf"):
                pdf_files.append(os.path.join(root, file))
    return pdf_files

def getCandidatesInfos(folder):
    candidates_infos = []
    pdf_files = getPdfFiles(folder)
    infos2extract = [
        {"odyssee": "Nom de famille",
         "target": "Nom",
         "operation": (lambda x : x.upper())},
        {"odyssee": "Adresse électronique",
         "target": "Email"},
        {"odyssee": "Date de naissance",
         "target": "Date de naissance"},
        {"odyssee": "Lieu de naissance",
         "target": "Lieu de naissance"},
        {"odyssee": "Pays de nationalité",
         "target": "Pays de nationalité"},
        {"odyssee": "Téléphone personnel",
         "target": "Téléphone"},
        {"odyssee": "Adresse postale",
         "target": "Adresse postale"},
        {"odyssee": "Complément d’adresse",
         "target": "Adresse comp."},
        {"odyssee": "Code postal",
         "target": "Code postal"},
        {"odyssee": "Ville",
         "target": "Ville"},
        {"odyssee": "Pays",
         "target": "Pays"}
    ]
    for pdf_file in pdf_files:
        txt = getTextFromPDF(pdf_file)
        id = os.path.split(pdf_file)[-2].split("_")[-1]
        infos = extractInfos(infos2extract, txt)
        infos["id"] = id
        candidates_infos.append(infos)
    return candidates_infos

def stats(file):
    df = pd.read_excel(file)
    ages = []
    f = 0
    fav = []
    indecis = []
    defav = []

    for _, r in df.iterrows():
        ages.append(computeAge(r['Date de naissance']))
        if r['Civilité'] == 'Mme':
            f += 1
        if str(r['AvisRapp1']) not in ['Défavorable', 'Réservé', 'Favorable', 'Très favorable', 'nan'] or \
            str(r['AvisRapp2']) not in ['Défavorable', 'Réservé', 'Favorable', 'Très favorable', 'nan']:
            print("Probleme avis sur %s %s %s"%(r["Nom d'usage"], r['AvisRapp1'], r['AvisRapp2']))
            sys.exit(-1)
        if r['AvisRapp1'] in ['Favorable', 'Très favorable'] and r['AvisRapp2'] in ['Favorable', 'Très favorable']:
            fav.append("%s %s"%(r['Prénom'], r["Nom d'usage"]))
        elif r['AvisRapp1'] in ['Défavorable', 'Réservé'] and r['AvisRapp2'] in ['Défavorable', 'Réservé']:
            defav.append("%s %s"%(r['Prénom'], r["Nom d'usage"]))
        else:
            indecis.append("%s %s"%(r['Prénom'], r["Nom d'usage"]))

    print("---------------------------")
    print("Dossiers fav ou très favorables des 2 rapporteurs : %s"%len(fav))
    print(fav)
    print("Dossiers indécis ou manque un avis %s"%len(indecis))
    print(indecis)    
    print("Dossiers defav / reservé des 2 rapporteurs : %s"%len(defav))
    print(defav)

    print("---------------------------")
    print("%s dossiers recevables"%len(ages))
    print("%s femmes (%2.1f%%), %s hommes (%2.1f%%)"%(f, f/len(ages)*100, len(ages)-f, (1-f/len(ages))*100))
    print("age moyen : %2.1f min : %2.1f, max : %2.1f"%(sum(ages)/len(ages), min(ages), max(ages)))

def getMembres(file):
    df = pd.read_excel(file, sheet_name="MembresComite")
    membres = []

    for _, r in df.iterrows():
        m = {}
        m['nom'] = r['Nom']
        m['prenom'] = r['Prénom']
        m['corps'] = r['Corps']
        m['section'] = r['Section CNU ou organisme']
        m['interne'] = r['Interne / Externe'] == 'Interne'
        m['etablissement'] = r['Etablissement'] if not pd.isna(r['Etablissement']) else ""
        if str(m['nom']) != "nan":
            m['nomfichier'] = unidecode("%s_%s"%(m['nom'].replace(" ","_"), \
                             m["prenom"].replace(" ","_")))
            membres.append(m)
    return membres

def generateImpartialite(file):
    membres = getMembres(file)
    saveJson(membres, 'membres.json')
    createDir('impartialite')
    os.system('node carbone/impartialite.js') 


avisdetailles = {
    ''    : ""
}

def loadAvisDetailles (file):
    df = pd.read_excel(file, sheet_name="AvisDetailles")

    for _, r in df.iterrows():
        avisdetailles[r['Code']] = r['Texte']

def checkAvis(avis):
    for a in avis:
        if a not in avisdetailles.keys():
            print("ERROR: avis " + a + " not found")

def getAvisTexte(avis):
    if avis == None:
        return ""
    s = ""
    for a in avis:
        s = s + avisdetailles[a] + " "
    return s

def formatAdresse(l):
    s = ""
    LB = "\n"
    if str(l['Adresse postale']) != "nan":
        s = s + l['Adresse postale'] + LB
    if str(l['Adresse comp.']) != "nan":
        s = "%s%s%s"%(s, l['Adresse comp.'], LB)

    if str(l['Pays']) == "FRANCE":
        s = "%s%s %s"%(s, int(l['Code postal']), l['Ville'])
    else:
        s = "%s%s %s%s%s"%(s, l['Code postal'], l['Ville'], LB, l['Pays'])
    return s

def getListAvis(r):
    a = r['AvisAudition']
    if str(a) != 'nan':
        return a.replace(" ", "").split(",")
    else:
        return []

def getCandidateData(r, rapNum):
    c = {}
    if (rapNum == 1):
        if str(r["Rapporteur1"]) == "nan" :
            rapporteur = "Rapporteur1"
        else:
            rapporteur = r["Rapporteur1"]
    else:
        if str(r["Rapporteur2"]) == "nan" :
            rapporteur = "Rapporteur2"
        else:
            rapporteur = r["Rapporteur2"]        

    listeAvis = getListAvis(r)
    checkAvis(listeAvis)
    if 'A' in listeAvis:
        avistexte = "Favorable"
        avisboolean = "true"
    else:
        avistexte = "Défavorable"
        avisboolean = "false"                
    avis = getAvisTexte(listeAvis)

    c["nomcandidat"] = r["Nom d'usage"]
    c["prenomcandidat"] = r["Prénom"]
    c["civilite"] = r["Civilité"]
    c["civiliteExt"] = "Madame" if r["Civilité"] == "Mme" else "Monsieur"
    c["heureAudition"] = formattxt(r["heureAudition"], "")
    c["adresse"] = formatAdresse(r)
    c["avisboolean"] = avisboolean
    c["avis"] = avis
    c["email"] = r['Email']
    c["avistexte"] = avistexte

    c["nomfichier"] = unidecode("%s_%s-%s"%(r["Nom d'usage"].replace(" ","_"), \
                         r["Prénom"].replace(" ","_"), rapporteur.replace(" ","_")))
    c["nomfichierCandidat"] = unidecode("%s_%s"%(r["Nom d'usage"].replace(" ","_"), \
                         r["Prénom"].replace(" ","_")))
    c["age"] = "%3.0f ans"%computeAge(r['Date de naissance'])
    c["titreThese"] = cleantext(r['Titre thèse'])
    c["DescriptionRecherche"] = cleantext(r['Description recherche'], True)
    c["MotsCles"] = cleantext(r['Mots clés'], True)
    c["LieuSoutenance"] = cleantext(r['Lieu soutenance'])
    c["DirecteurThese"] = cleantext(r['Directeur Thèse'])
    c["jury"] = cleantext(r['Jury'])
    s = rapporteur.split(" ")
    rapp = rapporteur
    if len(s) > 1:
        if len(s) == 3:
            rapp = "%s %s %s"%(s[2],s[0],s[1])
        else:
            rapp = "%s %s"%(s[1],s[0]) 
    c["rapporteur"] = rapp
    return c 

def generateReportsReunion1(file):
    df = pd.read_excel(file)
    candidats = []

    for _, r in df.iterrows():
        candidats.append(getCandidateData(r, 1))
        candidats.append(getCandidateData(r, 2))

    saveJson(candidats, 'candidats.json') 
    createDir('rapportsMembresReunion1')
    os.system('node carbone/generaterapportsMembresReunion1.js')

def findRapporteur(nomprenom, rapporteurs):
    if pd.isna(nomprenom):
        return None
    return next((r for r in rapporteurs if "%s %s"%(r["nomUsage"], r["prenom"]) == unidecode(nomprenom.upper())), None)

def designationRapporteurs(file, odyssee):
    df = pd.read_excel(file, sheet_name="Candidats")
    for _, row in df.iterrows():
        id_candidat = row["ID"]

        rapporteur1 = row["Rapporteur1"]
        rapp1 = findRapporteur(rapporteur1, odyssee.get_committee_members())
        if not pd.isna(rapporteur1) and not rapp1:
            print(f"Rapporteur {rapporteur1} non trouvé dans les rapporteurs")
            continue

        rapporteur2 = row["Rapporteur2"]
        rapp2 = findRapporteur(rapporteur2, odyssee.get_committee_members())
        if not pd.isna(rapporteur2) and not rapp2:
            print(f"Rapporteur {rapporteur2} non trouvé dans les rapporteurs")
            continue

        if not pd.isna(rapporteur1) and not pd.isna(rapporteur2):
            print(f"Assignation de {rapporteur1} et {rapporteur2} pour le candidat {row['Prénom']} {row["Nom d'usage"]}")
            res = odyssee.assign_jury_members_to_candidate(id_candidat, rapp1["id"], rapp2["id"])

            if res:
                print("OK")
            else:
                print("KO")
                sys.exit(1)

def getAvis(txt):
    tsLesAvis = [{"txt": "Très favorable à l’audition", "short": "TF"}, {"txt": "Favorable à l’audition", "short": "F"}, {"txt": "Réservé", "short": "R"}, {"txt": "Défavorable à l’audition", "short": "D"}]
    avis = []
    for a in [a["txt"] for a in tsLesAvis]:
        case = extractInfoRapport(a, txt)
        if "☒" in case or "X" in case:
            avis.append(a)
    if len(avis) == 0:
        return ""
    if len(avis) > 1:
        print("Attention : Plusieurs avis trouvés dans le rapport :", avis)
        return ""
    avis = next(a for a in tsLesAvis if a["txt"] == avis[0])
    return avis["short"]

def callbackOnReportsDownloaded(filename, res, r, i):
    txt = getTextFromPDF(filename)
    avis = getAvis(txt)
    rapporteur = f"{r['nomUsage']}_{r['prenom'].capitalize()}"
    res[f'rapporteur_{i+1}'] = {'nom': rapporteur, 'avis': avis}
    return res

def downloadReports(odyssee, path):
    allinfos = odyssee.download_reports(path)
    saveJson(allinfos, "infos_rapports.json")

    print(f"Total de rapports téléchargés: {allinfos['termines']} / {allinfos['total']}")

def reportRapporteursAvis(xlsxFile: str):
    """
    Report des avis des rapporteurs (de leur rapport pdf) dans le fichier Excel
    Args:        xlsxFile (str): chemin vers le fichier Excel 
    Returns:     Fichier Excel avec -2 dans le nom
    """
    candidates = loadJson("infos_rapports.json")
    candidates = candidates["infos"]

    wbIn = load_workbook(filename = xlsxFile)
    sIn = wbIn['Candidats']
    columnsInfos = getColumnInfo(sIn)

    for i in range(2, sIn.max_row+1):
        index_ID = next((col['index'] for col in columnsInfos if col['column'] == "ID"), None)
        ID = sIn.cell(row=i, column=index_ID + 1).value
        candidat = next((c for c in candidates if c["candidatID"] == ID ), None)
        if candidat is not None:
            rapport1 = candidat.get("rapport_1", None)
            if rapport1 is not None:
                txt = getTextFromPDF(rapport1)
                avis = getAvis(txt)
                colAvis1 = next((col['index'] for col in columnsInfos if col['column'] == "AvisRapp1"), None)
                sIn.cell(row=i, column=colAvis1 + 1).value = trAvis(avis)

            rapport2 = candidat.get("rapport_2", None)
            if rapport2 is not None:
                txt = getTextFromPDF(rapport2)
                avis = getAvis(txt)
                colAvis2 = next((col['index'] for col in columnsInfos if col['column'] == "AvisRapp2"), None)
                sIn.cell(row=i, column=colAvis2 + 1).value = trAvis(avis)

    newFile = xlsxFile.replace(".xlsx", "-2.xlsx")
    wbIn.save(newFile)
    print(f"Avis des rapporteurs reportés dans le fichier {newFile}")

def avisCandidatEtVotePremiereReunion(xlsxFile, odyssee):
    avisdetailles = loadAvisDetailles(xlsxFile)

    df = pd.read_excel(xlsxFile, sheet_name="Candidats")
    for _, r in tqdm.tqdm(df.iterrows(), desc="Enregistrement des avis de la première réunion"):
        id_candidat = r["ID"]
        avis = getListAvis(r)
        cleaned_avis = checkAvis(avis)
        motif = getAvisTexte(cleaned_avis)
        nombreVotants, suffragesExprimes, bulletinsNuls, bulletinsBlancs, bulletinsEnAccord, bulletinsEnDesaccord = r["nombreVotants"], r["suffragesExprimes"], r["bulletinsNuls"], r["bulletinsBlancs"], r["bulletinsEnAccord"], r["bulletinsEnDesaccord"]
        odyssee.opinion_for_interview(id_candidat, avis, motif, nombreVotants, suffragesExprimes, bulletinsNuls, bulletinsBlancs, bulletinsEnAccord, bulletinsEnDesaccord)

def generatePVrepartition(file):
    df = pd.read_excel(file)
    candidats = []

    for _, r in df.iterrows():
        c = getCandidateData(r, 1)
        c2 = getCandidateData(r, 2)
        c['nomrap1'] = c['rapporteur']
        c['nomrap2'] = c2['rapporteur']
        c['nomprenom'] = "%s - %s"%(c['nomcandidat'], c['prenomcandidat'])
        candidats.append(c)

    candidats = sorted(candidats, key=lambda x: x['nomprenom'])
    saveJson(candidats, 'candidats.json')
    os.system('node carbone/generaterapportsPVrepartition.js')   

def generateAttestationsVisio(file):
    membres = getMembres(file)
    saveJson(membres, 'membres.json')
    createDir('attestationsVisio')
    os.system('node carbone/attestationsVisio.js')  

def generateLettresAuditionnes(file):
    df = pd.read_excel(file)
    candidats = []

    for _, r in df.iterrows():
        c = getCandidateData(r, 1)
        if c["avisboolean"] == "true":
            candidats.append(c)

    saveJson(candidats, 'candidats.json')
    createDir('convocationsCandiats')
    os.system('node carbone/generateLettresAuditionnes.js')   

def sendMailAuditionnes(file, config):
    df = pd.read_excel(file)
    candidats = []

    for _, r in df.iterrows():
        c = getCandidateData(r, 1)
        if c["avistexte"] == "Favorable":
            candidats.append(c)

    candidats = candidats[0:1]

    for c in candidats:
        subject = "Convocation à l'audition du poste XXXX"
        recipient = "%s %s"%(c["prenomcandidat"], c["nomcandidat"])
        recipient = recipient.lower().title()
        c['recipient'] = recipient
        print('%s/%s.pdf'%(config['FILES']['convocations-candidats'], c['nomfichierCandidat']))
        # Read the content of templates/mailCandodatsAuditionnes.txt
        with open('templates/mailCandidatsAuditionnes.txt', 'r') as f:
            mail_template = f.read()
        message = mail_template.format_map(c)
        thunderbird(config['SOFTWARE']['thunderbird-bin'], recipientName=recipient, recipientAddress=c['email'], thesubject=subject, thecontent=message, cc='', attachment='%s/%s.pdf'%(config['FILES']['convocations-candidats'], c['nomfichierCandidat']))


def main():
    config = configparser.ConfigParser()
    config.read("config.ini")

    parser = argparse.ArgumentParser(description='ezCOS - un outil pour faciliter la gestion des comités de sélection')
    parser.add_argument('-auth', help = 'Authentification sur Odyssee', action="store_true")
    parser.add_argument('-downloadApplications', help = 'Téléchargement des dossiers des candidats', action="store_true")
    parser.add_argument('-unzipApplications', help = 'Décompression des dossiers des candidats', action="store_true")
    parser.add_argument('-populateXLSX', help = 'Récupération des infos sur les candidats sur Odyssee', action="store_true")
    parser.add_argument('-stats', help = 'Compute stats', action="store_true")
    parser.add_argument('-impartialite', help = "Déclaration d'impartialité", action="store_true")
    parser.add_argument('-r1', help = 'rapports précomplétés sur les candidats', action="store_true")
    parser.add_argument('-assignOdyssee', help = 'affectation des dossiers aux rapporteurs sur Odyssee', action="store_true")
    parser.add_argument('-pvRepart', help = 'generate PV repartition', action="store_true")
    parser.add_argument('-attestationsVisio', help = 'attestions participation visio membres', action="store_true")
    parser.add_argument('-reports', help = "téléchargement des rapports des candidatures d'Odyssee", action="store_true")
    parser.add_argument('-reports2xlsx', help = 'report des avis des rapporteurs dans le fichier Excel', action="store_true")
    parser.add_argument('-avisPremiereReunion', help = 'avis par candidat + vote suite à la première réunion', action="store_true")
    parser.add_argument('-convocCandidats', help = 'generate convocations candidats', action="store_true")
    parser.add_argument('-convocCandidatsPDF', help = 'generate convocations candidats', action="store_true")
    parser.add_argument('-envoimailauditionnes', help = 'Envoi de mail aux candidats auditionnés', action="store_true")
    args = parser.parse_args()

    numposte = config["ODYSSEE"]["numposte"]
    odyssee = OdysseeCOS(numposte)

    candidates_folder = "dossiers_candidats"

    xlsx_candidats = config['FILES']['xlsx-candidats']
    soffice = config['SOFTWARE']['soffice']

    if args.auth:
        login = config["ODYSSEE"]["login"]
        password = config["ODYSSEE"]["password"]
        odyssee.authenticate(login, password)

    if args.downloadApplications:
        odyssee.download_applications(candidates_folder)

    if args.unzipApplications:
        unzip_applications(candidates_folder)

    if args.populateXLSX:
        template_xlsx = "comite.xlsx"
        populateXLSX('templates', template_xlsx, odyssee, candidates_folder)

    if args.stats:
        stats(xlsx_candidats)

    if args.impartialite:
        generateImpartialite(xlsx_candidats)

    if args.r1:
        loadAvisDetailles(xlsx_candidats)
        generateReportsReunion1(xlsx_candidats)

    if args.assignOdyssee:
        designationRapporteurs(xlsx_candidats, odyssee)

    if args.pvRepart:
        loadAvisDetailles(xlsx_candidats)
        generatePVrepartition(xlsx_candidats)

    if args.attestationsVisio:
        generateAttestationsVisio(xlsx_candidats)

    if args.reports:
        loadAvisDetailles(xlsx_candidats)
        downloadReports(odyssee, "rapportsOdyssee")

    if args.reports2xlsx:
        reportRapporteursAvis(config["FILES"]["xlsx-candidats"])

    if args.avisPremiereReunion:
        avisCandidatEtVotePremiereReunion(config["FILES"]["xlsx-candidats"],  odyssee)

    if args.convocCandidats:
        loadAvisDetailles(xlsx_candidats)
        generateLettresAuditionnes(xlsx_candidats)

    if args.convocCandidatsPDF:
        os.system('cd convocationsCandiats; %s --headless --convert-to pdf *.docx'%soffice)

    if args.envoimailauditionnes:
        loadAvisDetailles(xlsx_candidats)
        sendMailAuditionnes(xlsx_candidats, config)

if __name__ == "__main__":
    main()